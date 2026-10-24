from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

doc = Document()
doc.sections[0].header.paragraphs[0].text = 'Northstar synthetic header'
doc.sections[0].footer.paragraphs[0].text = 'Synthetic footer — internal example'
doc.add_heading('Synthetic launch document', 0)
doc.add_paragraph('A document fixture for editor testing. No real company material.')
p = doc.add_paragraph()
r = p.add_run('Cloudberry Daily Cleanser costs INR 799 per 150ml.')
r.bold = True
r.font.name = 'Calibri'
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string('167569')
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
table.cell(0,0).text = 'Product'
table.cell(0,1).text = 'Approved price'
table.cell(1,0).text = 'Cloudberry Daily Cleanser'
table.cell(1,1).text = 'INR 799'
doc.add_paragraph('Gentle daily care.', style='List Bullet')
doc.add_paragraph('Results vary. Read the product label before use.')
path = Path('.local/test-results/editor-fixture.docx')
path.parent.mkdir(parents=True,exist_ok=True)
doc.save(path)
print(path)
