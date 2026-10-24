import json
from pathlib import Path
from docx import Document

root=Path('.local/test-results')
original=Document(root/'editor-fixture.docx')
edited=Document(root/'editor-roundtrip.docx')
checks=[]
def check(name,condition):
    checks.append({'name':name,'passed':bool(condition)})
    if not condition:raise AssertionError(name)
try:
    check('Header text retained',original.sections[0].header.paragraphs[0].text==edited.sections[0].header.paragraphs[0].text)
    check('Footer text retained',original.sections[0].footer.paragraphs[0].text==edited.sections[0].footer.paragraphs[0].text)
    check('Table shape retained',len(edited.tables)==1 and len(edited.tables[0].rows)==2 and len(edited.tables[0].columns)==2)
    check('Edited table cell retained',edited.tables[0].cell(1,1).text=='INR 899')
    check('Paragraph edit retained',any('Tested edit.' in p.text for p in edited.paragraphs))
    p=next(p for p in edited.paragraphs if p.text.startswith('Cloudberry Daily Cleanser costs'))
    check('Untouched bold run retained',all(r.bold for r in p.runs if r.text.strip()))
    check('Untouched font size retained',all(r.font.size and r.font.size.pt==13 for r in p.runs if r.text.strip()))
    check('Untouched color retained',all(str(r.font.color.rgb)=='167569' for r in p.runs if r.text.strip()))
    check('Section page size retained',edited.sections[0].page_width==original.sections[0].page_width and edited.sections[0].page_height==original.sections[0].page_height)
    status='passed'
except Exception as error:
    status='failed';checks.append({'error':str(error)})
report={'kind':'Synthetic DOCX roundtrip fidelity subset','status':status,'checks':checks,'limitations':'This representative fixture does not prove full Word parity or arbitrary DOCX fidelity.'}
(root/'docx-fidelity.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2))
raise SystemExit(0 if status=='passed' else 1)
