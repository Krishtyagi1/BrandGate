"""Runs in a bounded child process; uploaded text is data, never instructions."""

import json
import sys
import zipfile
from pathlib import Path


def extract(path):
    sections = []
    if path.suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are unsupported. Upload an unlocked text-based PDF.")
        if len(reader.pages) > 100:
            raise ValueError("PDFs must contain no more than 100 pages.")
        for number, page in enumerate(reader.pages, 1):
            sections.append({"location": f"Page {number}", "text": page.extract_text() or ""})
    else:
        with zipfile.ZipFile(path) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024
            ):
                raise ValueError("DOCX expands beyond the supported extraction limit.")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("This file is not a valid DOCX document.")
        from docx import Document

        document = Document(path)
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                sections.append({"location": f"Paragraph {number}", "text": paragraph.text.strip()})
        for number, table in enumerate(document.tables, 1):
            sections.append(
                {
                    "location": f"Table {number}",
                    "text": "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows),
                }
            )
    text = "\n\n".join(section["text"] for section in sections)
    if len(text) > 200000:
        raise ValueError("Extracted text exceeds 200,000 characters. Split this document.")
    if not text.strip():
        raise ValueError(
            "No selectable text was found. Scanned files need OCR, which is not available in this phase."
        )
    partial = any(not section["text"].strip() for section in sections)
    return {
        "text": text,
        "locations": sections,
        "status": "partial" if partial else "extracted",
    }


if __name__ == "__main__":
    try:
        print(json.dumps(extract(Path(sys.argv[1])), ensure_ascii=True))
    except Exception as error:
        print(json.dumps({"error": str(error)[:300]}))
        sys.exit(1)
