from io import BytesIO

from conftest import add_member
from docx import Document


def test_draft_immutable_versions_and_stale_save(workspace):
    owner, org, _ = workspace
    base = f"/api/organisations/{org}/drafts"
    created = owner.post(
        base, json={"title": "Synthetic launch", "channel": "website", "text": "First text"}
    )
    assert created.status_code == 201, created.text
    draft = created.json()
    first = draft["current_version_id"]
    body = {"expected_current_version_id": first, "text": "Updated text"}
    saved = owner.post(f"{base}/{draft['id']}/versions", json=body)
    assert saved.status_code == 201, saved.text
    assert owner.post(f"{base}/{draft['id']}/versions", json=body).status_code == 409
    history = owner.get(f"{base}/{draft['id']}").json()["versions"]
    assert len(history) == 2
    assert history[-1]["text"] == "First text"


def test_creator_ownership_and_cross_tenant_visibility(workspace, account):
    owner, org, _ = workspace
    creator, email, _ = account()
    add_member(owner, org, creator, email)
    base = f"/api/organisations/{org}/drafts"
    draft = owner.post(base, json={"title": "Private draft", "text": "Owner text"}).json()
    assert creator.get(f"{base}/{draft['id']}").status_code == 404
    assert creator.get(base).json() == []
    outsider, _, _ = account()
    assert outsider.get(f"{base}/{draft['id']}").status_code == 404


def test_unicode_docx_versions_retain_original_bytes(workspace):
    owner, org, _ = workspace
    doc = Document()
    doc.add_paragraph("Cleanser — ₹799 🪷")
    doc.sections[0].header.paragraphs[0].text = "Header — private"
    stream = BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    base = f"/api/organisations/{org}/drafts"
    uploaded = owner.post(
        f"{base}/upload",
        files={
            "file": (
                "example.docx",
                data,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    draft = uploaded.json()
    version = draft["current_version"]
    assert "₹799 🪷" in version["text"]
    assert "Header — private" in version["text"]
    assert owner.get(f"{base}/{draft['id']}/versions/{version['id']}/download").content == data
    doc.paragraphs[0].text = "Cleanser — ₹899 🪷"
    stream = BytesIO()
    doc.save(stream)
    saved = owner.post(
        f"{base}/{draft['id']}/file-versions",
        data={"expected_current_version_id": version["id"]},
        files={
            "file": (
                "example.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert saved.status_code == 201, saved.text
    assert saved.json()["artifact_sha256"] != version["artifact_sha256"]
    assert owner.get(f"{base}/{draft['id']}/versions/{version['id']}/download").content == data
