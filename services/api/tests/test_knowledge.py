"""Synthetic source collection and immutable confirmation regressions."""

from importlib.util import find_spec
from io import BytesIO

from conftest import add_member
from docx import Document


def test_knowledge_backend_exists():
    assert find_spec("app.knowledge") is not None


def document(text):
    result = BytesIO()
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(result)
    return result.getvalue()


def setup_view(owner, org):
    response = owner.post(f"/api/organisations/{org}/knowledge/new-version")
    assert response.status_code == 201, response.text
    return response.json()


def proposal(owner, org, version, **changes):
    response = owner.post(
        f"/api/organisations/{org}/knowledge/proposals",
        json={
            "version_id": version,
            "type": "tone",
            "rule_key": "voice",
            "value": "Friendly",
            "check": "tone",
            **changes,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def confirm(owner, org, view, ids, replacements=None):
    return owner.post(
        f"/api/organisations/{org}/knowledge/confirm",
        json={
            "version_id": view["version"]["id"],
            "proposal_ids": ids,
            "replace_rule_ids": replacements or {},
            "expected_snapshot_id": view["snapshot_id"],
        },
    )


def test_knowledge_confirmation_snapshots_and_explicit_replacement(workspace):
    owner, org, _ = workspace
    view = setup_view(owner, org)
    assert view["checks"] == [] and view["rules"] == [] and view["sources"] == []
    first = proposal(owner, org, view["version"]["id"])
    response = confirm(owner, org, view, [first["id"]])
    assert response.status_code == 200, response.text
    active = response.json()
    assert active["version"]["status"] == "active"
    assert active["snapshot_id"] and len(active["rules"]) == 1
    change = proposal(owner, org, view["version"]["id"], value="Professional")
    assert change["existing_rule_id"] == active["rules"][0]["id"]
    assert confirm(owner, org, active, [change["id"]]).status_code == 409
    updated = confirm(
        owner, org, active, [change["id"]], {change["id"]: change["existing_rule_id"]}
    )
    assert updated.status_code == 200, updated.text
    latest = updated.json()
    assert latest["version"]["id"] == active["version"]["id"]
    assert latest["snapshot_id"] != active["snapshot_id"]
    assert [r["value"] for r in latest["rules"]] == ["Professional"]
    assert confirm(owner, org, active, []).status_code == 409
    empty = setup_view(owner, org)
    assert empty["sources"] == [] and empty["rules"] == [] and empty["checks"] == []


def test_knowledge_exact_evidence_and_aggregate_gaps(workspace):
    owner, org, _ = workspace
    view = setup_view(owner, org)
    version = view["version"]["id"]
    for text in ["Tone: Friendly and clear.", "Required wording: Terms apply."]:
        uploaded = owner.post(
            f"/api/organisations/{org}/knowledge/sources",
            data={"version_id": version},
            files={
                "file": (
                    "rules.docx",
                    document(text),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 201, uploaded.text
    saved = owner.post(
        f"/api/organisations/{org}/knowledge/preferences",
        json={
            "version_id": version,
            "checks": ["tone", "required_wording", "prices"],
            "instructions": "Keep claims clear.",
        },
    )
    assert saved.status_code == 200, saved.text
    current = saved.json()
    source = current["sources"][0]
    forged = owner.post(
        f"/api/organisations/{org}/knowledge/proposals",
        json={
            "version_id": version,
            "type": "tone",
            "rule_key": "tone",
            "value": "Friendly",
            "check": "tone",
            "source_document_id": source["id"],
            "source_excerpt": "Never in the file",
            "source_location": "Paragraph 1",
        },
    )
    assert forged.status_code == 422
    assert all(s["download_url"] and s["locations"] for s in current["sources"])
    assert all(item["status"] == "not_checked" for item in current["coverage"])


def test_knowledge_owner_and_tenant_guards(workspace, account):
    owner, org, _ = workspace
    other, _, _ = account()
    assert other.get(f"/api/organisations/{org}/knowledge").status_code == 404
    creator, email, _ = account()
    add_member(owner, org, creator, email)
    assert creator.post(f"/api/organisations/{org}/knowledge/new-version").status_code == 403


def test_knowledge_ai_excerpts_are_literal_and_all_sections_processed():
    from ingestion.knowledge_intake import chunks, validate_candidate

    source = {
        "id": "00000000-0000-0000-0000-000000000001",
        "locations": [
            {"location": f"Paragraph {n}", "text": f"Tone: literal line {n}"} for n in range(1, 42)
        ],
    }
    parts = list(chunks([source]))
    assert sum(len(p) for p in parts) == 41
    good = {
        "type": "tone",
        "rule_key": "tone",
        "value": "literal",
        "check": "tone",
        "source_document_id": source["id"],
        "source_excerpt": "literal line 41",
        "source_location": "Paragraph 41",
        "fact": None,
    }
    assert validate_candidate(good, [source])["source_excerpt"] == "literal line 41"
    good["source_excerpt"] = "literal line 99"
    import pytest

    with pytest.raises(ValueError, match="literal"):
        validate_candidate(good, [source])


def test_knowledge_aggregate_analysis_reports_missing_and_unconfirmed(workspace, monkeypatch):
    from ingestion import knowledge_intake
    from providers.openrouter import POLICY_VERSION

    owner, org, _ = workspace
    view = setup_view(owner, org)
    version = view["version"]["id"]
    owner.post(
        f"/api/organisations/{org}/knowledge/preferences",
        json={
            "version_id": version,
            "checks": ["tone", "required_wording", "prices"],
            "instructions": "",
        },
    )
    for text in ["Tone: Friendly and clear.", "Required wording: Terms apply."]:
        assert (
            owner.post(
                f"/api/organisations/{org}/knowledge/sources",
                data={"version_id": version},
                files={
                    "file": (
                        "rules.docx",
                        document(text),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            ).status_code
            == 201
        )
    assert (
        owner.post(
            f"/api/organisations/{org}/knowledge/analyse", json={"version_id": version}
        ).status_code
        == 409
    )
    assert (
        owner.post(
            f"/api/organisations/{org}/ai/consent",
            json={"policy_version": POLICY_VERSION, "enabled": True},
        ).status_code
        == 200
    )

    def synthetic_provider(system, payload, schema, purpose, usage_callback):
        proposals = []
        for item in payload["sources"]:
            tone = item["text"].startswith("Tone:")
            proposals.append(
                {
                    "type": "tone" if tone else "required_disclaimer",
                    "rule_key": "voice" if tone else "terms",
                    "value": item["text"],
                    "check": "tone" if tone else "required_wording",
                    "source_document_id": item["source_document_id"],
                    "source_location": item["source_location"],
                    "source_excerpt": item["text"],
                    "fact": None,
                }
            )
        return schema.model_validate({"proposals": proposals}), {"model": "synthetic-test-only"}

    monkeypatch.setattr(knowledge_intake, "complete", synthetic_provider)
    assert (
        owner.post(
            f"/api/organisations/{org}/knowledge/analyse", json={"version_id": version}
        ).status_code
        == 202
    )
    current = owner.get(
        f"/api/organisations/{org}/knowledge", params={"version_id": version}
    ).json()
    assert current["extraction"]["status"] == "completed"
    assert len(current["proposals"]) == 2
    assert current["rules"] == [] and current["snapshot_id"] is None
    assert {c["check"]: c["status"] for c in current["coverage"]} == {
        "tone": "needs_confirmation",
        "required_wording": "needs_confirmation",
        "prices": "source_missing",
    }
    approved = confirm(owner, org, current, [p["id"] for p in current["proposals"]])
    assert approved.status_code == 200, approved.text
    assert [c["status"] for c in approved.json()["coverage"]] == [
        "available",
        "available",
        "source_missing",
    ]
    # A new unread document changes aggregate coverage; absence is no longer proven.
    owner.post(
        f"/api/organisations/{org}/knowledge/sources",
        data={"version_id": version},
        files={
            "file": (
                "added.docx",
                document("Price: INR 150."),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    fresh = owner.get(f"/api/organisations/{org}/knowledge", params={"version_id": version}).json()
    assert fresh["coverage"][-1]["status"] == "not_checked"


def test_knowledge_image_reference_is_human_confirmed_and_snapshot_pinned(workspace):
    from PIL import Image

    from app.db import transaction
    from app.knowledge_snapshots import snapshot_references, snapshot_rules

    owner, org, _ = workspace
    view = setup_view(owner, org)
    image = BytesIO()
    Image.new("RGB", (8, 8), "green").save(image, format="PNG")
    uploaded = owner.post(
        f"/api/organisations/{org}/knowledge/sources",
        data={"version_id": view["version"]["id"]},
        files={"file": ("logo.png", image.getvalue(), "image/png")},
    )
    assert uploaded.status_code == 201, uploaded.text
    staged = uploaded.json()
    assert staged["rules"] == [] and staged["sources"][0]["extraction_status"] == "reference_only"
    approved = confirm(owner, org, staged, [staged["proposals"][0]["id"]])
    assert approved.status_code == 200, approved.text
    active = approved.json()
    user = owner.get("/api/auth/me").json()
    with transaction(user, org) as conn:
        refs = snapshot_references(conn, org, active["snapshot_id"])
        assert len(refs) == 1 and refs[0]["source_location"] == "Reference image"
        assert refs[0]["media_type"] == "image/png" and refs[0]["sha256"]
        assert len(snapshot_rules(conn, org, active["snapshot_id"])) == 1
    preference = owner.post(
        f"/api/organisations/{org}/knowledge/preferences",
        json={
            "version_id": active["version"]["id"],
            "checks": ["logos"],
            "instructions": "Compare logos",
        },
    ).json()
    assert preference["snapshot_id"] != active["snapshot_id"]
    with transaction(user, org) as conn:
        assert snapshot_references(conn, org, active["snapshot_id"]) == refs


def test_knowledge_docx_embedded_images_are_private_staged_references(workspace):
    from PIL import Image

    owner, org, _ = workspace
    view = setup_view(owner, org)
    picture = BytesIO()
    Image.new("RGB", (12, 12), "blue").save(picture, format="PNG")
    picture.seek(0)
    doc = Document()
    doc.add_paragraph("Tone: clear and friendly.")
    doc.add_picture(picture)
    data = BytesIO()
    doc.save(data)
    uploaded = owner.post(
        f"/api/organisations/{org}/knowledge/sources",
        data={"version_id": view["version"]["id"]},
        files={
            "file": (
                "master.docx",
                data.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    current = uploaded.json()
    images = [s for s in current["sources"] if s["media_type"].startswith("image/")]
    assert len(images) == 1
    assert "word/media/" in images[0]["locations"][0]["location"]
    assert current["snapshot_id"] is None
    assert current["proposals"][0]["check"] == "logos"
    downloaded = owner.get(images[0]["download_url"])
    assert downloaded.status_code == 200 and downloaded.headers["content-type"] == "image/png"


def test_knowledge_price_overlap_and_source_value_validation(workspace):
    import pytest

    from app.knowledge_snapshots import rules_overlap
    from ingestion.knowledge_intake import validate_candidate

    base = {
        "type": "approved_fact",
        "rule_key": "price",
        "scope_channel": "global",
        "fact": {
            "kind": "price",
            "entity": "Serum",
            "amount": "100",
            "currency": "INR",
            "unit": "bottle",
            "effective_from": "2026-01-01",
            "effective_to": None,
        },
    }
    changed = {
        **base,
        "rule_key": "new_price",
        "fact": {**base["fact"], "amount": "150", "effective_from": "2026-10-24"},
    }
    assert rules_overlap(base, changed)
    historical = {**base, "fact": {**base["fact"], "effective_to": "2026-10-23"}}
    assert not rules_overlap(historical, changed)
    source = {
        "id": "00000000-0000-0000-0000-000000000001",
        "locations": [
            {"location": "Paragraph 1", "text": "Serum INR 100 per bottle, effective 2026-01-01."}
        ],
    }
    candidate = {
        **base,
        "value": "Serum INR 100 per bottle",
        "check": "prices",
        "source_document_id": source["id"],
        "source_excerpt": source["locations"][0]["text"],
        "source_location": "Paragraph 1",
    }
    assert validate_candidate(candidate, [source])["fact"]["amount"] == "100"
    with pytest.raises(ValueError, match="amount and currency"):
        validate_candidate({**candidate, "fact": {**base["fact"], "amount": "150"}}, [source])
    with pytest.raises(ValueError, match="wording"):
        validate_candidate({**candidate, "value": "Serum INR 150 per bottle"}, [source])
    owner, org, _ = workspace
    view = setup_view(owner, org)
    first = proposal(
        owner,
        org,
        view["version"]["id"],
        type="approved_fact",
        rule_key="price",
        value="Serum INR 100",
        check="prices",
        fact=base["fact"],
    )
    active = confirm(owner, org, view, [first["id"]]).json()
    update = proposal(
        owner,
        org,
        view["version"]["id"],
        type="approved_fact",
        rule_key="different_key",
        value="Serum INR 150",
        check="prices",
        fact=changed["fact"],
    )
    assert update["existing_rule_id"] == active["rules"][0]["id"]
    assert confirm(owner, org, active, [update["id"]]).status_code == 409


def test_knowledge_provider_failure_keeps_sources_and_never_claims_missing(workspace, monkeypatch):
    from ingestion import knowledge_intake
    from providers.openrouter import POLICY_VERSION, ProviderUnavailable

    owner, org, _ = workspace
    view = setup_view(owner, org)
    version = view["version"]["id"]
    owner.post(
        f"/api/organisations/{org}/knowledge/preferences",
        json={"version_id": version, "checks": ["prices"], "instructions": ""},
    )
    owner.post(
        f"/api/organisations/{org}/knowledge/sources",
        data={"version_id": version},
        files={
            "file": (
                "facts.docx",
                document("Serum INR 100 per bottle."),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    owner.post(
        f"/api/organisations/{org}/ai/consent",
        json={"policy_version": POLICY_VERSION, "enabled": True},
    )

    def unavailable(*args, **kwargs):
        raise ProviderUnavailable(
            "Free model rate limit reached. Retry after the provider quota resets."
        )

    monkeypatch.setattr(knowledge_intake, "complete", unavailable)
    assert (
        owner.post(
            f"/api/organisations/{org}/knowledge/analyse", json={"version_id": version}
        ).status_code
        == 202
    )
    current = owner.get(
        f"/api/organisations/{org}/knowledge", params={"version_id": version}
    ).json()
    assert current["extraction"]["status"] == "partial"
    assert "rate limit" in current["extraction"]["error"]
    assert current["coverage"][0]["status"] == "not_checked"
    assert len(current["sources"]) == 1 and current["rules"] == []


def test_knowledge_manual_price_forgery_and_pending_edits_require_confirmation(workspace):
    owner, org, _ = workspace
    view = setup_view(owner, org)
    version = view["version"]["id"]
    excerpt = "Serum INR 100 per bottle, effective 2026-01-01."
    uploaded = owner.post(
        f"/api/organisations/{org}/knowledge/sources",
        data={"version_id": version},
        files={
            "file": (
                "facts.docx",
                document(excerpt),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    body = {
        "version_id": version,
        "type": "approved_fact",
        "rule_key": "serum_price",
        "value": "Serum INR 150 per bottle",
        "check": "prices",
        "source_document_id": uploaded["sources"][0]["id"],
        "source_excerpt": excerpt,
        "source_location": "Paragraph 1",
        "fact": {
            "kind": "price",
            "entity": "Serum",
            "amount": "150",
            "currency": "INR",
            "unit": "bottle",
            "effective_from": "2026-01-01",
            "effective_to": None,
        },
    }
    created = owner.post(f"/api/organisations/{org}/knowledge/proposals", json=body)
    assert created.status_code == 422
    body["value"] = "Serum INR 100 per bottle"
    body["fact"]["amount"] = "100"
    created = owner.post(f"/api/organisations/{org}/knowledge/proposals", json=body)
    assert created.status_code == 201, created.text
    candidate = created.json()
    url = f"/api/organisations/{org}/knowledge/proposals/{candidate['id']}"
    edited = owner.patch(url, json={**body, "rule_key": "daily_serum_price"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["status"] == "pending" and edited.json()["rule_key"] == "daily_serum_price"
    assert (
        owner.get(f"/api/organisations/{org}/knowledge", params={"version_id": version}).json()[
            "snapshot_id"
        ]
        is None
    )
    assert (
        owner.patch(url, json={**body, "fact": {**body["fact"], "amount": "150"}}).status_code
        == 422
    )
    other = setup_view(owner, org)
    assert owner.patch(url, json={**body, "version_id": other["version"]["id"]}).status_code == 404
    approved = confirm(owner, org, view, [candidate["id"]])
    assert approved.status_code == 200, approved.text
    assert owner.patch(url, json=body).status_code == 409


def test_knowledge_redacted_price_fields_restore_before_validation(monkeypatch):
    from ingestion import knowledge_intake

    source = {
        "id": "00000000-0000-0000-0000-000000000001",
        "locations": [
            {
                "location": "Paragraph 1",
                "text": "Serum INR 100 per bottle, effective 2026-01-01. Contact owner@example.com.",
            }
        ],
    }
    batch = list(knowledge_intake.chunks([source]))[0]

    def synthetic_provider(system, payload, schema, purpose, usage_callback):
        item = payload["sources"][0]
        assert "2026-01-01" not in item["text"]
        assert "owner@example.com" not in item["text"]
        return schema.model_validate(
            {
                "proposals": [
                    {
                        "type": "approved_fact",
                        "rule_key": "serum_price",
                        "value": "Serum INR 100 per bottle",
                        "check": "prices",
                        "source_document_id": item["source_document_id"],
                        "source_location": item["source_location"],
                        "source_excerpt": item["text"],
                        "fact": {
                            "kind": "price",
                            "entity": "Serum",
                            "amount": "100",
                            "currency": "INR",
                            "unit": "bottle",
                            "effective_from": "[PRIVATE_1]",
                            "effective_to": None,
                        },
                    }
                ]
            }
        ), {"model": "synthetic-test-only"}

    monkeypatch.setattr(knowledge_intake, "complete", synthetic_provider)
    candidates, _, invalid = knowledge_intake.analyse_chunk(batch, [source])
    assert invalid == 0 and len(candidates) == 1
    assert candidates[0]["fact"]["effective_from"] == "2026-01-01"
    assert candidates[0]["source_excerpt"] == source["locations"][0]["text"]


def test_knowledge_pending_legacy_rule_can_be_edited_without_activating(workspace):
    owner, org, _ = workspace
    view = setup_view(owner, org)
    version = view["version"]["id"]
    created = owner.post(
        f"/api/organisations/{org}/guidelines/versions/{version}/rules",
        json={
            "type": "tone",
            "rule_key": "price_discussion_tone",
            "value": "Friendly",
            "approved": False,
        },
    )
    assert created.status_code == 201, created.text
    candidate = created.json()
    current = owner.get(
        f"/api/organisations/{org}/knowledge", params={"version_id": version}
    ).json()
    assert current["proposals"][0]["check"] == "tone"
    updated = owner.patch(
        f"/api/organisations/{org}/knowledge/proposals/{candidate['id']}",
        json={
            "version_id": version,
            "type": "tone",
            "rule_key": "price_discussion_tone",
            "value": "Calm",
            "check": "tone",
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["status"] == "pending"
    current = owner.get(
        f"/api/organisations/{org}/knowledge", params={"version_id": version}
    ).json()
    assert len(current["proposals"]) == 1 and current["proposals"][0]["value"] == "Calm"
    assert current["rules"] == []
